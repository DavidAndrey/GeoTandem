"""Read an import file into a uniform ``Source`` (F-2.1, F-2.2).

Vector formats go through GDAL (pyogrio's raw API, no geopandas); CSV through
the standard library; Excel through openpyxl. Nothing here touches the data
core: a ``Source`` is what the file says, before any decision about it.
"""

from __future__ import annotations

import csv
import io
import logging
from collections.abc import Sequence
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Literal

import numpy as np
import openpyxl
import pyogrio
import pyogrio.raw
import shapely
from pydantic import BaseModel, Field
from pyproj import CRS
from pyproj.exceptions import CRSError
from shapely.geometry.base import BaseGeometry

from geotandem import gdal
from geotandem.data.interface import AttributeType
from geotandem.importing.values import as_date, infer

gdal.check(pyogrio.list_drivers(read=True))

Format = Literal["geojson", "shapefile", "gpkg", "csv", "xlsx"]

FORMATS: dict[str, Format] = {
    ".geojson": "geojson",
    ".json": "geojson",
    ".zip": "shapefile",
    ".gpkg": "gpkg",
    ".csv": "csv",
    ".txt": "csv",
    ".xlsx": "xlsx",
}
VECTOR_FORMATS: frozenset[Format] = frozenset({"geojson", "shapefile", "gpkg"})
FORMAT_NAMES: dict[Format, str] = {
    "geojson": "GeoJSON",
    "shapefile": "zipped shapefile",
    "gpkg": "GeoPackage",
    "csv": "CSV",
    "xlsx": "Excel workbook",
}

log = logging.getLogger(__name__)
DELIMITERS = ",;\t|"
_SNIFF_BYTES = 64 * 1024


class ReadOptions(BaseModel):
    """How to read a file; every field is detected when left empty."""

    sublayer: str | None = None
    """GeoPackage layer, shapefile inside the zip, or Excel sheet."""
    encoding: str | None = None
    """Text encoding of a CSV file or a shapefile without ``.cpg``."""
    delimiter: str | None = Field(default=None, min_length=1, max_length=1)
    """CSV field delimiter, one character."""


class Message(BaseModel):
    """A finding about the source; ``code`` is stable, ``message`` explains it."""

    code: str
    message: str
    column: str | None = None
    count: int | None = None


def _unreadable(file_name: str, fmt: Format, exc: Exception) -> SourceError:
    """The reason in the server log, not in the answer: GDAL and openpyxl name the
    staged file's path on the server and their own internals (security review #10)."""
    log.warning("import file %r unreadable as %s: %s", file_name, fmt, exc)
    return SourceError(
        "unreadable",
        f"'{file_name}' cannot be read as a {FORMAT_NAMES[fmt]}. "
        "It may be damaged or of another format.",
        file=file_name,
    )


class SourceError(Exception):
    """The file cannot be read as given (design D11: blocks the wizard)."""

    def __init__(self, code: str, message: str, **details: Any) -> None:
        super().__init__(message)
        self.code = code
        self.message = message
        self.details = details


@dataclass
class Column:
    name: str
    """As written in the source."""
    data_type: AttributeType
    values: list[Any]
    """Converted to ``data_type``; ``None`` for empty cells."""


@dataclass
class Source:
    format: Format
    file_name: str
    options: ReadOptions
    """The options actually used, detected ones filled in."""
    sublayers: list[str]
    columns: list[Column]
    geometries: list[BaseGeometry | None] | None
    """``None`` for tabular sources."""
    crs: int | None = None
    """EPSG code of the geometries, if the file names an identifiable one."""
    crs_label: str | None = None
    """What the file says about its CRS, for display."""
    notes: list[Message] = field(default_factory=list)

    @property
    def record_count(self) -> int:
        if self.geometries is not None:
            return len(self.geometries)
        return len(self.columns[0].values) if self.columns else 0

    @property
    def is_vector(self) -> bool:
        return self.geometries is not None

    def column(self, name: str) -> Column:
        for column in self.columns:
            if column.name == name:
                return column
        raise KeyError(name)


def detect_format(file_name: str) -> Format:
    suffix = Path(file_name).suffix.lower()
    if suffix not in FORMATS:
        raise SourceError(
            "unsupported_format",
            f"'{suffix or file_name}' is not a supported format. Supported: GeoJSON, "
            "Shapefile (as .zip), GeoPackage, CSV, Excel (.xlsx).",
            file=file_name,
        )
    return FORMATS[suffix]


def read_source(path: Path, file_name: str, options: ReadOptions | None = None) -> Source:
    options = options or ReadOptions()
    fmt = detect_format(file_name)
    if fmt in VECTOR_FORMATS:
        return _read_vector(path, file_name, fmt, options)
    if fmt == "csv":
        return _read_csv(path, file_name, options)
    return _read_xlsx(path, file_name, options)


# --- vector -------------------------------------------------------------------

_OGR_TYPES: dict[str, AttributeType] = {
    "OFTInteger": "integer",
    "OFTInteger64": "integer",
    "OFTReal": "real",
    "OFTString": "text",
}
_OGR_DATES = {"OFTDate", "OFTDateTime", "OFTTime"}


def _read_vector(path: Path, file_name: str, fmt: Format, options: ReadOptions) -> Source:
    target = str(path)
    try:
        sublayers = [str(name) for name, _ in pyogrio.list_layers(target)]
    except pyogrio.errors.DataSourceError as exc:
        raise _unreadable(file_name, fmt, exc) from exc
    if not sublayers:
        raise SourceError("no_layers", "The file contains no layer.")
    sublayer = options.sublayer or sublayers[0]
    if sublayer not in sublayers:
        raise SourceError(
            "unknown_sublayer", f"The file has no layer '{sublayer}'.", sublayers=sublayers
        )
    try:
        meta, _, wkb, fields = pyogrio.raw.read(target, layer=sublayer, encoding=options.encoding)
    except (pyogrio.errors.DataSourceError, pyogrio.errors.DataLayerError) as exc:
        raise _unreadable(file_name, fmt, exc) from exc

    notes: list[Message] = []
    columns = []
    names, repeated = _distinct([str(name) for name in meta["fields"]])
    if repeated:
        notes.append(_duplicate_header(repeated))
    for name, ogr_type, subtype, data in zip(
        names, meta["ogr_types"], meta["ogr_subtypes"], fields, strict=True
    ):
        data_type, values, as_text = _ogr_column(ogr_type, subtype, data)
        columns.append(Column(str(name), data_type, values))
        if as_text:
            notes.append(
                Message(
                    code="stored_as_text",
                    message=f"Column '{name}' ({ogr_type}) is stored as text.",
                    column=str(name),
                )
            )
    if wkb is not None:
        geometries: list[BaseGeometry | None] = list(shapely.from_wkb(wkb))
    else:  # a layer without geometry column, e.g. a GeoPackage attribute table
        geometries = [None] * (len(fields[0]) if len(fields) else 0)
    crs, label = _crs(meta.get("crs"))
    return Source(
        format=fmt,
        file_name=file_name,
        options=ReadOptions(sublayer=sublayer, encoding=meta.get("encoding")),
        sublayers=sublayers,
        columns=columns,
        geometries=geometries,
        crs=crs,
        crs_label=label,
        notes=notes,
    )


def _ogr_column(
    ogr_type: str, subtype: str, data: np.ndarray
) -> tuple[AttributeType, list[Any], bool]:
    """Type from the OGR field definition: numpy turns integers with nulls into floats."""
    if ogr_type in _OGR_DATES:
        texts = np.datetime_as_string(data) if data.dtype.kind == "M" else data.astype(str)
        values = [None if t == "NaT" else str(t) for t in texts]
        dates = [None if v is None else as_date(v.replace("/", "-")[:10]) for v in values]
        if ogr_type == "OFTDate" and all(
            d is not None for d, v in zip(dates, values, strict=True) if v
        ):
            return "date", dates, False
        return "text", values, True
    if subtype == "OFSTBoolean":
        return "boolean", [None if _null(v) else bool(v) for v in data], False
    data_type = _OGR_TYPES.get(ogr_type)
    if data_type == "integer":
        return "integer", [None if _null(v) else int(v) for v in data], False
    if data_type == "real":
        return "real", [None if _null(v) else float(v) for v in data], False
    if data_type == "text":
        return "text", [None if v is None else str(v) for v in data], False
    return "text", [None if _null(v) else str(v) for v in data], True


def _null(value: Any) -> bool:
    return value is None or (isinstance(value, float | np.floating) and np.isnan(value))


def _crs(text: str | None) -> tuple[int | None, str | None]:
    if not text:
        return None, None
    try:
        crs = CRS.from_user_input(text)
    except CRSError:
        return None, text
    return crs.to_epsg(min_confidence=70), crs.name


# --- tabular ------------------------------------------------------------------


def _table(
    fmt: Format,
    file_name: str,
    options: ReadOptions,
    sublayers: list[str],
    rows: Sequence[Sequence[Any]],
    notes: list[Message],
) -> Source:
    """Header row plus data rows → typed columns."""
    rows = [r for r in rows if any(c is not None and str(c).strip() for c in r)]
    if not rows:
        raise SourceError("empty", "The file contains no rows.")
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    header = [h or f"spalte_{i + 1}" for i, h in enumerate(header)]
    header, repeated = _distinct(header)
    if repeated:
        notes.append(_duplicate_header(repeated))
    body = rows[1:]
    ragged = sum(1 for r in body if len(r) != len(header))
    if ragged:
        notes.append(
            Message(
                code="ragged_rows",
                message=f"{ragged} rows have a different number of fields than the header.",
                count=ragged,
            )
        )
    if len(header) == 1 and fmt == "csv":
        notes.append(
            Message(
                code="single_column",
                message="Only one column was found; the delimiter may be wrong.",
            )
        )
    columns = []
    for i, name in enumerate(header):
        data_type, values, dates = infer([r[i] if i < len(r) else None for r in body])
        columns.append(Column(name, data_type, values))
        if dates:
            notes.append(
                Message(
                    code="stored_as_text",
                    message=f"Column '{name}' contains times of day; they are stored as text.",
                    column=name,
                )
            )
    return Source(
        format=fmt,
        file_name=file_name,
        options=options,
        sublayers=sublayers,
        columns=columns,
        geometries=None,
        notes=notes,
    )


def _distinct(header: list[str]) -> tuple[list[str], list[str]]:
    """Each column a name of its own, the file's where it can: "Name", "Name (2)".

    Decisions name a column by it, so two alike would be one column twice.
    """
    taken = set(header)
    seen: set[str] = set()
    names, repeated = [], []
    for name in header:
        unique, n = name, 1
        # Not a name the file gives another column further on either.
        while unique in seen or (unique != name and unique in taken):
            n += 1
            unique = f"{name} ({n})"
        seen.add(unique)
        names.append(unique)
        if unique != name:
            repeated.append(name)
    return names, sorted(set(repeated))


def _duplicate_header(repeated: list[str]) -> Message:
    return Message(
        code="duplicate_header",
        message=f"Columns named alike are told apart by a number: {', '.join(repeated)}.",
        count=len(repeated),
    )


def _read_csv(path: Path, file_name: str, options: ReadOptions) -> Source:
    raw = path.read_bytes()
    notes: list[Message] = []
    encoding = options.encoding
    if encoding is None:
        try:
            raw.decode("utf-8-sig")
            encoding = "utf-8-sig"
        except UnicodeDecodeError:
            encoding = "cp1252"
            notes.append(
                Message(
                    code="encoding_fallback",
                    message="The file is not UTF-8; it was read as Windows-1252.",
                )
            )
    try:
        text = raw.decode(encoding)
    except (UnicodeDecodeError, LookupError) as exc:
        raise SourceError(
            "unreadable_encoding", f"The file cannot be decoded as {encoding}.", encoding=encoding
        ) from exc
    delimiter = options.delimiter or _sniff(text[:_SNIFF_BYTES])
    rows = list(csv.reader(io.StringIO(text, newline=""), delimiter=delimiter))
    resolved = ReadOptions(encoding=encoding, delimiter=delimiter)
    return _table("csv", file_name, resolved, [], rows, notes)


def _sniff(sample: str) -> str:
    try:
        return csv.Sniffer().sniff(sample, delimiters=DELIMITERS).delimiter
    except csv.Error:
        header = sample.splitlines()[0] if sample else ""
        return max(DELIMITERS, key=header.count) if header else ","


def _read_xlsx(path: Path, file_name: str, options: ReadOptions) -> Source:
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises a zoo of exception types
        raise _unreadable(file_name, "xlsx", exc) from exc
    try:
        sheets = list(workbook.sheetnames)
        sheet = options.sublayer or sheets[0]
        if sheet not in sheets:
            raise SourceError(
                "unknown_sublayer", f"The workbook has no sheet '{sheet}'.", sublayers=sheets
            )
        rows = list(workbook[sheet].iter_rows(values_only=True))
    finally:
        workbook.close()
    return _table("xlsx", file_name, ReadOptions(sublayer=sheet), sheets, rows, [])
